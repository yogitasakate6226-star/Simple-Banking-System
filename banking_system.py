from openpyxl import Workbook, load_workbook
import os
file_name = "banking_system.xlsx"

if os.path.exists(file_name):
    wb = load_workbook(file_name)
    ws = wb.active
else:
    wb = Workbook()
    ws = wb.active
    ws.append(["Account No", "Name", "Balance"])
    wb.save(file_name)
    

import openpyxl
import os

# Create or open Excel file
if os.path.exists("banking_system.xlsx"):
    wb = openpyxl.load_workbook("banking_system.xlsx")
    sheet = wb.active
else:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Account No", "Name", "Mobile", "PIN", "Balance"])
    wb.save("banking_system.xlsx")


# Create Account
def create_account():
    name = input("Enter your name: ")
    mobile = input("Enter mobile number: ")
    pin = input("Enter 4-digit PIN: ")
    balance = float(input("Enter initial deposit: "))

    if len(mobile) != 10 or not mobile.isdigit():
        print("Invalid mobile number")
        return

    if len(pin) != 4 or not pin.isdigit():
        print("Invalid PIN")
        return

    if balance < 0:
        print("Invalid amount")
        return

    acc_no = sheet.max_row + 1000

    sheet.append([acc_no, name, mobile, pin, balance])
    wb.save("banking_system.xlsx")

    print("Account created successfully!")
    print("Your account number:", acc_no)


# Find Account
def find_account(acc_no):
    for row in range(2, sheet.max_row + 1):
        if str(sheet.cell(row, 1).value) == acc_no:
            return row
    return None


# Deposit Money
def deposit_money():
    acc_no = input("Enter account number: ")
    row = find_account(acc_no)

    if row is None:
        print("Account not found")
        return

    pin = input("Enter PIN: ")

    if pin != str(sheet.cell(row, 4).value):
        print("Wrong PIN")
        return

    amount = float(input("Enter deposit amount: "))

    if amount <= 0:
        print("Invalid amount")
        return

    balance = sheet.cell(row, 5).value
    balance = balance + amount
    sheet.cell(row, 5).value = balance

    wb.save("banking_system.xlsx")
    print("Deposit successful!")
    print("Current balance:", balance)


# Withdraw Money
def withdraw_money():
    acc_no = input("Enter account number: ")
    row = find_account(acc_no)

    if row is None:
        print("Account not found")
        return

    pin = input("Enter PIN: ")

    if pin != str(sheet.cell(row, 4).value):
        print("Wrong PIN")
        return

    amount = float(input("Enter withdrawal amount: "))

    if amount <= 0:
        print("Invalid amount")
        return

    balance = sheet.cell(row, 5).value

    if amount > balance:
        print("Insufficient balance")
    else:
        balance = balance - amount
        sheet.cell(row, 5).value = balance

        wb.save("banking_system.xlsx")
        print("Withdrawal successful!")
        print("Current balance:", balance)


# Check Balance
def check_balance():
    acc_no = input("Enter account number: ")
    row = find_account(acc_no)

    if row is None:
        print("Account not found")
        return

    pin = input("Enter PIN: ")

    if pin != str(sheet.cell(row, 4).value):
        print("Wrong PIN")
        return

    print("Account Number:", acc_no)
    print("Name:", sheet.cell(row, 2).value)
    print("Balance:", sheet.cell(row, 5).value)


# Main Menu
def main():
    while True:
        print("\n===== BANKING SYSTEM =====")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Check Balance")
        print("5. Exit")

        choice = input("Enter your choice: ")

        if choice == "1":
            create_account()

        elif choice == "2":
            deposit_money()

        elif choice == "3":
            withdraw_money()

        elif choice == "4":
            check_balance()

        elif choice == "5":
            print("Thank you!")
            break

        else:
            print("Invalid choice")


main()
